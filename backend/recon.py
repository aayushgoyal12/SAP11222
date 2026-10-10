import io
import re
import uuid
from collections import Counter
from datetime import datetime, date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MAX_ROWS = 50000
GSTIN_RE = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]$")
COMPONENT_KEYS = ("igst", "cgst", "sgst")
DISCLAIMER = ("Review assistance only. This report flags records for human review; it is not tax or legal advice "
              "and nothing is posted to SAP or the GST portal.")

CANONICAL = [
    {"key": "gstin", "label": "Supplier GSTIN", "required": True,
     "aliases": ["gstin", "gstin of supplier", "supplier gstin", "lictradnum", "federal tax id", "gstin/uin",
                 "vendor gstin", "gst number", "gst no", "gstin no", "tax id"]},
    {"key": "invoice_no", "label": "Invoice Number", "required": True,
     "aliases": ["invoice number", "invoice no", "numatcard", "supplier invoice no", "vendor ref no",
                 "vendor reference", "bill no", "inv no", "invoice num", "xblnr", "reference"]},
    {"key": "invoice_date", "label": "Invoice Date", "required": True,
     "aliases": ["invoice date", "taxdate", "document date", "bill date", "inv date", "bldat", "doc date"]},
    {"key": "taxable_value", "label": "Taxable Value", "required": True,
     "aliases": ["taxable value", "taxable amount", "baseamount", "base amount", "assessable value", "taxable amt"]},
    {"key": "tax_amount", "label": "Total Tax Amount", "required": False,
     "aliases": ["tax amount", "vatsum", "total tax", "total tax amount", "gst amount", "tax amt"]},
    {"key": "igst", "label": "IGST", "required": False, "aliases": ["igst", "integrated tax", "igst amount"]},
    {"key": "cgst", "label": "CGST", "required": False, "aliases": ["cgst", "central tax", "cgst amount"]},
    {"key": "sgst", "label": "SGST / UTGST", "required": False,
     "aliases": ["sgst", "state/ut tax", "state tax", "utgst", "sgst amount"]},
    {"key": "vendor_name", "label": "Vendor Name", "required": False,
     "aliases": ["cardname", "vendor name", "supplier name", "trade/legal name", "trade name", "legal name",
                 "party name", "name1"]},
]
LABELS = {f["key"]: f["label"] for f in CANONICAL}
DATE_MODES = {
    "AUTO": "auto-detect (ISO/YYYYMMDD first, then day-first DD-MM-YYYY)",
    "DMY": "DD-MM-YYYY (day first)",
    "MDY": "MM-DD-YYYY (month first)",
    "YMD": "YYYY-MM-DD / YYYYMMDD",
}
DATE_FORMATS = {
    "DMY": ["%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y", "%d.%m.%y", "%d-%b-%Y", "%d-%b-%y",
            "%d %b %Y", "%d/%b/%Y"],
    "MDY": ["%m-%d-%Y", "%m/%d/%Y", "%m.%d.%Y", "%m/%d/%y", "%b %d %Y", "%b %d, %Y"],
    "YMD": ["%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"],
}
DATE_ORDER = {"AUTO": ["YMD", "DMY"], "DMY": ["YMD", "DMY"], "MDY": ["YMD", "MDY"], "YMD": ["YMD"]}


def norm_header(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


ALIAS_NORMS = {f["key"]: [norm_header(a) for a in f["aliases"]] for f in CANONICAL}
ALL_ALIASES = {a for v in ALIAS_NORMS.values() for a in v}


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, (pd.Timestamp, datetime)):
        return "" if pd.isna(v) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _score(row):
    return sum(1 for c in row if c and norm_header(c) in ALL_ALIASES)


def _parse_sheet(raw):
    rows = [[_cell(v) for v in r] for r in raw.itertuples(index=False, name=None)]
    rows = [r for r in rows]
    if not rows:
        return None
    scores = [_score(r) for r in rows[:25]]
    best = max(range(len(scores)), key=lambda i: scores[i])
    if scores[best] == 0:
        best = next((i for i, r in enumerate(rows) if any(r)), 0)
    header = rows[best]
    start = best + 1
    if best + 1 < len(rows) and _score(rows[best + 1]) >= 2:
        header = [b or a for a, b in zip(header, rows[best + 1])]
        start = best + 2
    return rows, header, best, start, scores[best]


def read_table(content, filename):
    name = (filename or "").lower()
    if not content:
        raise ValueError("File is empty.")
    try:
        if name.endswith((".csv", ".txt")):
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("latin-1")
            sheets = {"CSV": pd.read_csv(io.StringIO(text), header=None, dtype=str, keep_default_na=False,
                                         sep=None, engine="python")}
        elif name.endswith((".xlsx", ".xlsm")):
            sheets = pd.read_excel(io.BytesIO(content), header=None, dtype=object, engine="openpyxl", sheet_name=None)
        else:
            raise ValueError("Unsupported file type. Upload a CSV or XLSX export.")
    except ValueError as e:
        if "Unsupported" in str(e):
            raise
        raise ValueError(f"Could not read file: {e}")
    except Exception as e:
        raise ValueError(f"Could not read file ({type(e).__name__}). Check it is a valid, unprotected CSV/XLSX.")

    parsed = [(sn, _parse_sheet(df)) for sn, df in sheets.items()]
    parsed = [(sn, p) for sn, p in parsed if p]
    if not parsed:
        raise ValueError("No rows found in file.")
    sheet, (rows, header, best, start, _) = max(parsed, key=lambda x: x[1][4])

    width = max(len(header), max((len(r) for r in rows[start:]), default=0))
    header = list(header) + [""] * (width - len(header))
    data, nums = [], []
    for i in range(start, len(rows)):
        r = list(rows[i]) + [""] * (width - len(rows[i]))
        if any(r):
            data.append(r)
            nums.append(i + 1)
    if not data:
        raise ValueError("Header row found but no data rows below it.")
    if len(data) > MAX_ROWS:
        raise ValueError(f"File has {len(data)} rows; the MVP limit is {MAX_ROWS}.")

    keep = [j for j in range(width) if header[j] or any(r[j] for r in data)]
    cols, seen = [], {}
    for j in keep:
        h = header[j] or f"Column {j + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        cols.append(h)
    data = [[r[j] for j in keep] for r in data]
    return {"columns": cols, "rows": data, "row_nums": nums, "header_row": best + 1, "sheet": str(sheet)}


def auto_map(columns):
    mapping, confidence, used = {}, {}, set()
    normed = {c: norm_header(c) for c in columns}
    for f in CANONICAL:
        key = f["key"]
        mapping[key], confidence[key] = None, "none"
        for alias in ALIAS_NORMS[key]:
            hit = next((c for c in columns if c not in used and normed[c] == alias), None)
            if hit:
                mapping[key], confidence[key] = hit, "high"
                used.add(hit)
                break
        if mapping[key]:
            continue
        for alias in ALIAS_NORMS[key]:
            if len(alias) < 5:
                continue
            hit = next((c for c in columns if c not in used and alias in normed[c]), None)
            if hit:
                mapping[key], confidence[key] = hit, "medium"
                used.add(hit)
                break
    return mapping, confidence


def validate_mapping(mapping, columns, side_label):
    errors = []
    for f in CANONICAL:
        col = mapping.get(f["key"])
        if col and col not in columns:
            errors.append(f"{side_label}: column '{col}' for {f['label']} not found in file")
        if f["required"] and not col:
            errors.append(f"{side_label}: {f['label']} must be mapped")
    if not mapping.get("tax_amount") and not any(mapping.get(k) for k in COMPONENT_KEYS):
        errors.append(f"{side_label}: map Total Tax Amount or at least one of IGST/CGST/SGST")
    return errors


def norm_inv(s):
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper()).lstrip("0")


def parse_date(s, mode):
    s = (s or "").strip()
    if not s:
        return None
    if re.match(r"^\d{4}-\d{2}-\d{2}[ T]", s):
        s = s[:10]
    for key in DATE_ORDER.get(mode, DATE_ORDER["AUTO"]):
        for f in DATE_FORMATS[key]:
            try:
                return datetime.strptime(s, f).date()
            except ValueError:
                pass
    return None


def _amount(raw):
    s = (raw or "").replace(",", "").replace("₹", "").replace("INR", "").replace("Rs.", "").replace(" ", "")
    if not s:
        return None, None
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    if s.endswith("-"):
        neg, s = True, s[:-1]
    try:
        v = float(s)
    except ValueError:
        return None, f"'{raw}' is not a number"
    return round(-v if neg else v, 2), None


def normalize_side(columns, rows, row_nums, mapping, date_mode):
    idx = {k: columns.index(c) for k, c in mapping.items() if c and c in columns}
    out = []
    for r, rn in zip(rows, row_nums):
        def get(k):
            return r[idx[k]].strip() if k in idx and idx[k] < len(r) else ""
        errors = []
        gstin = re.sub(r"\s", "", get("gstin").upper())
        if not gstin:
            errors.append("GSTIN is blank")
        elif not GSTIN_RE.match(gstin):
            errors.append(f"GSTIN '{gstin}' is not a valid 15-character GSTIN format")
        inv_raw = get("invoice_no")
        inv = norm_inv(inv_raw)
        if not inv:
            errors.append("Invoice number is blank")
        d_raw = get("invoice_date")
        d = parse_date(d_raw, date_mode)
        if not d_raw:
            errors.append("Invoice date is blank")
        elif d is None:
            errors.append(f"Invoice date '{d_raw}' could not be parsed as a valid date")
        taxable, err = _amount(get("taxable_value"))
        if err:
            errors.append(f"Taxable value {err}")
        elif taxable is None:
            errors.append("Taxable value is blank")
        if "tax_amount" in idx:
            tax, err = _amount(get("tax_amount"))
            if err:
                errors.append(f"Tax amount {err}")
            elif tax is None:
                errors.append("Tax amount is blank")
        else:
            parts = []
            for k in COMPONENT_KEYS:
                if k in idx:
                    v, err = _amount(get(k))
                    if err:
                        errors.append(f"{k.upper()} {err}")
                    elif v is not None:
                        parts.append(v)
            tax = round(sum(parts), 2)
        out.append({"row": rn, "gstin": gstin, "inv": inv, "inv_raw": inv_raw, "date": d, "date_raw": d_raw,
                    "taxable": taxable, "tax": tax, "vendor": get("vendor_name"), "errors": errors})
    return out


def _diff(a, b):
    return round(a - b, 2) if a is not None and b is not None else None


def _res(status, reason, s=None, c=None, notes=None):
    def d(x):
        if not x:
            return None
        return x["date"].isoformat() if x["date"] else x["date_raw"]
    st, ct = (s or {}).get("taxable"), (c or {}).get("taxable")
    sx, cx = (s or {}).get("tax"), (c or {}).get("tax")
    return {
        "status": status, "reason": reason, "notes": notes or [],
        "gstin": (s or c)["gstin"], "vendor_name": (s or {}).get("vendor") or (c or {}).get("vendor") or "",
        "sap_invoice_no": s["inv_raw"] if s else None, "cmp_invoice_no": c["inv_raw"] if c else None,
        "sap_date": d(s), "cmp_date": d(c),
        "sap_taxable": st, "cmp_taxable": ct, "taxable_diff": _diff(st, ct),
        "sap_tax": sx, "cmp_tax": cx, "tax_diff": _diff(sx, cx),
        "sap_row": s["row"] if s else None, "cmp_row": c["row"] if c else None,
        "reviewed": False, "review_note": "",
    }


def _money(v):
    return f"₹{v:,.2f}"


def reconcile(sap, cmp, sap_map, cmp_map, tol, sap_mode, cmp_mode, cmp_label):
    S = normalize_side(sap["columns"], sap["rows"], sap["row_nums"], sap_map, sap_mode)
    C = normalize_side(cmp["columns"], cmp["rows"], cmp["row_nums"], cmp_map, cmp_mode)
    results = []

    cmp_index = {}
    for c in C:
        if c["errors"]:
            results.append(_res("INVALID", f"{cmp_label} row {c['row']}: " + "; ".join(c["errors"]), c=c))
            continue
        key = (c["gstin"], c["inv"])
        if key in cmp_index:
            results.append(_res("DUPLICATE", f"Duplicate in {cmp_label}: same GSTIN + invoice number as "
                                             f"{cmp_label} row {cmp_index[key]['row']}", c=c))
            continue
        cmp_index[key] = c

    seen, used, missing = {}, set(), []
    for s in S:
        if s["errors"]:
            results.append(_res("INVALID", f"SAP row {s['row']}: " + "; ".join(s["errors"]), s=s))
            continue
        key = (s["gstin"], s["inv"])
        if key in seen:
            results.append(_res("DUPLICATE", f"Duplicate in SAP export: same GSTIN + invoice number as SAP row "
                                             f"{seen[key]['row']} (possible double posting)", s=s))
            continue
        seen[key] = s
        c = cmp_index.get(key)
        if not c:
            r = _res("MISSING_IN_COMPARISON",
                     f"No {cmp_label} record with GSTIN {s['gstin']} and invoice no '{s['inv_raw']}'. Supplier may not "
                     f"have reported it yet, or the invoice number/GSTIN differs.", s=s)
            results.append(r)
            missing.append((r, s))
            continue
        used.add(key)
        dt, dx = _diff(s["taxable"], c["taxable"]), _diff(s["tax"], c["tax"])
        notes = []
        if s["inv_raw"].strip().upper() != c["inv_raw"].strip().upper():
            notes.append(f"Invoice number matched after normalisation ('{s['inv_raw']}' vs '{c['inv_raw']}')")
        issues = []
        if abs(dt) > tol:
            issues.append(f"Taxable value differs by {_money(dt)} (tolerance {_money(tol)})")
        if abs(dx) > tol:
            issues.append(f"Tax amount differs by {_money(dx)} (tolerance {_money(tol)})")
        date_gap = (s["date"] - c["date"]).days
        if issues:
            if date_gap:
                issues.append(f"invoice date also differs by {abs(date_gap)} day(s)")
            results.append(_res("AMOUNT_MISMATCH", "; ".join(issues), s, c, notes))
            continue
        if dt or dx:
            notes.append(f"Within tolerance: taxable diff {_money(dt)}, tax diff {_money(dx)}")
        if date_gap:
            results.append(_res("DATE_MISMATCH", f"Amounts agree but invoice date differs: SAP {s['date']} vs "
                                                 f"{cmp_label} {c['date']} ({abs(date_gap)} day(s))", s, c, notes))
        else:
            results.append(_res("MATCHED", "GSTIN, invoice number, date and amounts agree", s, c, notes))

    unmatched_cmp = [c for k, c in cmp_index.items() if k not in used]
    for c in unmatched_cmp:
        results.append(_res("MISSING_IN_SAP", f"Present in {cmp_label} but not found in SAP export (not booked yet, "
                                              f"or booked with a different invoice number/GSTIN)", c=c))
    for r, s in missing:
        hint = next((c for c in unmatched_cmp if c["gstin"] == s["gstin"]
                     and abs(c["taxable"] - s["taxable"]) <= tol), None)
        if hint:
            r["notes"].append(f"Possible match: {cmp_label} invoice '{hint['inv_raw']}' (row {hint['row']}) has the "
                              f"same GSTIN and taxable value — check the invoice number")

    for i, r in enumerate(results):
        r["seq"] = i + 1
        r["id"] = str(uuid.uuid4())

    counts = Counter(r["status"] for r in results)
    valid_s = [s for s in S if not s["errors"]]
    summary = {
        "total_records": len(S), "cmp_records": len(C),
        "matched": counts["MATCHED"], "missing_in_cmp": counts["MISSING_IN_COMPARISON"],
        "missing_in_sap": counts["MISSING_IN_SAP"], "amount_mismatch": counts["AMOUNT_MISMATCH"],
        "date_mismatch": counts["DATE_MISMATCH"], "duplicates": counts["DUPLICATE"], "invalid": counts["INVALID"],
        "requires_review": len(results) - counts["MATCHED"], "result_rows": len(results),
        "match_rate": round(counts["MATCHED"] / len(S) * 100, 1) if S else 0,
        "sap_taxable_total": round(sum(s["taxable"] for s in valid_s), 2),
        "sap_tax_total": round(sum(s["tax"] for s in valid_s), 2),
    }
    assumptions = [
        "Records are paired on Supplier GSTIN + normalised invoice number; date and amounts are then compared.",
        "Invoice numbers are compared case-insensitively after removing spaces, slashes, hyphens and leading zeros.",
        f"Amount tolerance of {_money(tol)} is applied separately to taxable value and tax amount.",
        f"SAP dates parsed as {DATE_MODES[sap_mode]}; {cmp_label} dates parsed as {DATE_MODES[cmp_mode]}.",
        "Duplicates: the first occurrence is reconciled; later occurrences are flagged for review.",
        "Cess, place of supply, reverse charge and ITC eligibility are not evaluated.",
    ]
    for label, m in (("SAP", sap_map), (cmp_label, cmp_map)):
        if not m.get("tax_amount"):
            assumptions.append(f"{label} tax amount computed as IGST + CGST + SGST (blank components treated as 0).")
    return results, summary, assumptions


STATUS_LABELS = {
    "MATCHED": "Matched", "MISSING_IN_COMPARISON": "Missing in {cmp}", "MISSING_IN_SAP": "Missing in SAP",
    "AMOUNT_MISMATCH": "Amount mismatch", "DATE_MISMATCH": "Date mismatch", "DUPLICATE": "Duplicate",
    "INVALID": "Invalid record",
}
STATUS_FILLS = {
    "MATCHED": "D1FAE5", "MISSING_IN_COMPARISON": "FEE2E2", "MISSING_IN_SAP": "FEF3C7", "AMOUNT_MISMATCH": "EDE9FE",
    "DATE_MISMATCH": "E0F2FE", "DUPLICATE": "FCE7F3", "INVALID": "E2E8F0",
}
HEADER_FILL = PatternFill("solid", fgColor="0F172A")


def _label(status, cmp_label):
    return STATUS_LABELS[status].format(cmp=cmp_label)


def _write_sheet(wb, title, results, cmp_label):
    ws = wb.create_sheet(title[:31])
    headers = ["#", "Status", "Reason flagged", "Notes / assumptions", "Supplier GSTIN", "Vendor",
               "SAP invoice no", f"{cmp_label} invoice no", "SAP date", f"{cmp_label} date",
               "SAP taxable", f"{cmp_label} taxable", "Taxable diff", "SAP tax", f"{cmp_label} tax", "Tax diff",
               "SAP row", f"{cmp_label} row", "Reviewed", "Review note"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in results:
        ws.append([r["seq"], _label(r["status"], cmp_label), r["reason"], " | ".join(r["notes"]), r["gstin"],
                   r["vendor_name"], r["sap_invoice_no"], r["cmp_invoice_no"], r["sap_date"], r["cmp_date"],
                   r["sap_taxable"], r["cmp_taxable"], r["taxable_diff"], r["sap_tax"], r["cmp_tax"], r["tax_diff"],
                   r["sap_row"], r["cmp_row"], "Yes" if r.get("reviewed") else "No", r.get("review_note") or ""])
        row = ws.max_row
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=STATUS_FILLS[r["status"]])
        for col in range(11, 17):
            ws.cell(row, col).number_format = "#,##0.00"
    widths = [6, 20, 60, 50, 19, 28, 20, 20, 12, 12, 14, 14, 13, 13, 13, 12, 8, 8, 9, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    if results:
        ws.auto_filter.ref = ws.dimensions


def build_report(job, results, exceptions_only):
    cmp_label = "Purchase Register" if job["comparison_type"] == "purchase_register" else "GSTR-2B"
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = job["summary"]
    ws.append(["SAP Export Reconciliation — " + ("Exception Report" if exceptions_only else "Full Reconciliation Report")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([DISCLAIMER])
    ws["A2"].font = Font(italic=True, color="B91C1C")
    ws.append([])
    meta = [
        ("Job", job["name"]), ("Generated (UTC)", datetime.utcnow().strftime("%Y-%m-%d %H:%M")),
        ("SAP export file", job["sap"]["filename"]), (f"{cmp_label} file", job["comparison"]["filename"]),
        ("Amount tolerance (₹)", job["settings"]["tolerance"]), ("Demo data", "Yes (fictional)" if job.get("is_demo") else "No"),
        ("", ""),
        ("SAP records", s["total_records"]), (f"{cmp_label} records", s["cmp_records"]), ("Matched", s["matched"]),
        (f"Missing in {cmp_label}", s["missing_in_cmp"]), ("Missing in SAP", s["missing_in_sap"]),
        ("Amount mismatches", s["amount_mismatch"]), ("Date mismatches", s["date_mismatch"]),
        ("Duplicates", s["duplicates"]), ("Invalid records", s["invalid"]),
        ("Records requiring review", s["requires_review"]), ("Match rate (% of SAP records)", s["match_rate"]),
    ]
    for k, v in meta:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(["Assumptions applied"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for a in job.get("assumptions", []):
        ws.append(["", a])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100

    exceptions = [r for r in results if r["status"] != "MATCHED"]
    if not exceptions_only:
        _write_sheet(wb, "All Records", results, cmp_label)
    _write_sheet(wb, "Exceptions", exceptions, cmp_label)
    for status in ["MISSING_IN_COMPARISON", "AMOUNT_MISMATCH", "DATE_MISMATCH", "DUPLICATE", "INVALID", "MISSING_IN_SAP"]:
        subset = [r for r in results if r["status"] == status]
        if subset:
            _write_sheet(wb, _label(status, cmp_label), subset, cmp_label)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
