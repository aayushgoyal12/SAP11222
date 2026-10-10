"""Backend API tests for SAP Export Reconciliation."""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ['REACT_APP_BACKEND_URL'].rstrip('/') if os.environ.get('REACT_APP_BACKEND_URL') else None
if not BASE_URL:
    # read frontend .env
    with open('/app/frontend/.env') as f:
        for line in f:
            if line.startswith('REACT_APP_BACKEND_URL='):
                BASE_URL = line.split('=', 1)[1].strip().rstrip('/')
                break

API = f"{BASE_URL}/api"

created_jobs = []
created_templates = []


@pytest.fixture(scope="session", autouse=True)
def cleanup():
    yield
    for jid in created_jobs:
        try:
            requests.delete(f"{API}/jobs/{jid}", timeout=15)
        except Exception:
            pass
    for tid in created_templates:
        try:
            requests.delete(f"{API}/templates/{tid}", timeout=15)
        except Exception:
            pass


# ---------- Basic ----------
def test_root():
    r = requests.get(f"{API}/", timeout=15)
    assert r.status_code == 200
    assert "message" in r.json()


def test_schema():
    r = requests.get(f"{API}/schema", timeout=15)
    assert r.status_code == 200
    data = r.json()
    assert "fields" in data and len(data["fields"]) >= 5
    assert "retention_days" in data
    assert "disclaimer" in data


# ---------- Demo files ----------
@pytest.mark.parametrize("kind", ["sap_csv", "sap_xlsx", "gstr2b_xlsx", "gstr2b_csv"])
def test_demo_file_download(kind):
    r = requests.get(f"{API}/demo/files/{kind}", timeout=30)
    assert r.status_code == 200
    assert len(r.content) > 100


def test_demo_file_unknown():
    r = requests.get(f"{API}/demo/files/unknown_kind", timeout=15)
    assert r.status_code == 404


# ---------- Demo job full flow (KPIs) ----------
@pytest.fixture(scope="module")
def demo_job():
    r = requests.post(f"{API}/jobs/demo", timeout=60)
    assert r.status_code == 200, r.text
    job = r.json()
    created_jobs.append(job["id"])
    return job


def test_demo_job_auto_mapping(demo_job):
    sap_map = demo_job["sap"]["suggested_mapping"]
    cmp_map = demo_job["comparison"]["suggested_mapping"]
    # SAP: GSTIN should map to LicTradNum
    assert sap_map["gstin"] == "LicTradNum", f"got {sap_map['gstin']}"
    assert sap_map["invoice_no"] == "NumAtCard"
    assert sap_map["taxable_value"] == "BaseAmount"
    # Comparison: IGST/CGST/SGST mapped, tax_amount should not be mapped
    assert cmp_map.get("igst") is not None
    assert cmp_map.get("cgst") is not None
    assert cmp_map.get("sgst") is not None
    assert cmp_map.get("tax_amount") is None, f"tax_amount should not be mapped, got {cmp_map.get('tax_amount')}"


def test_demo_job_gstr2b_header_detection(demo_job):
    # The GSTR-2B xlsx has title rows + 'Read me' sheet; header should be row 4 of B2B sheet
    assert demo_job["comparison"]["sheet"] == "B2B"
    assert demo_job["comparison"]["header_row"] == 4, f"got header_row={demo_job['comparison']['header_row']}"


def test_demo_reconcile_kpis(demo_job):
    body = {
        "sap_mapping": demo_job["sap"]["suggested_mapping"],
        "cmp_mapping": demo_job["comparison"]["suggested_mapping"],
        "tolerance": 1.0,
        "sap_date_format": "AUTO",
        "cmp_date_format": "AUTO",
    }
    r = requests.post(f"{API}/jobs/{demo_job['id']}/reconcile", json=body, timeout=60)
    assert r.status_code == 200, r.text
    s = r.json()["summary"]
    print("Summary:", s)
    assert s["total_records"] == 37, f"total {s['total_records']}"
    assert s["matched"] == 27, f"matched {s['matched']}"
    assert s["missing_in_cmp"] == 3, f"missing_in_cmp {s['missing_in_cmp']}"
    assert s["amount_mismatch"] == 3, f"amount_mismatch {s['amount_mismatch']}"
    assert s["duplicates"] == 1, f"duplicates {s['duplicates']}"
    assert s["requires_review"] == 14, f"requires_review {s['requires_review']}"


def test_demo_tolerance_zero_changes_matches():
    # Create fresh demo job and run with tolerance=0 -> matched 25
    r = requests.post(f"{API}/jobs/demo", timeout=60)
    assert r.status_code == 200
    job = r.json()
    created_jobs.append(job["id"])
    body = {
        "sap_mapping": job["sap"]["suggested_mapping"],
        "cmp_mapping": job["comparison"]["suggested_mapping"],
        "tolerance": 0.0,
        "sap_date_format": "AUTO",
        "cmp_date_format": "AUTO",
    }
    r = requests.post(f"{API}/jobs/{job['id']}/reconcile", json=body, timeout=60)
    assert r.status_code == 200
    s = r.json()["summary"]
    assert s["matched"] == 25, f"expected 25 matched with tol=0, got {s['matched']}"
    assert s["amount_mismatch"] == 5, f"amount_mismatch {s['amount_mismatch']}"


# ---------- Results & review ----------
def test_results_and_review(demo_job):
    r = requests.get(f"{API}/jobs/{demo_job['id']}/results", timeout=30)
    assert r.status_code == 200
    results = r.json()
    assert len(results) >= 37
    exc = next((x for x in results if x["status"] != "MATCHED"), None)
    assert exc is not None
    rid = exc["id"]
    r = requests.patch(f"{API}/jobs/{demo_job['id']}/results/{rid}",
                       json={"reviewed": True, "review_note": "Checked - OK"}, timeout=30)
    assert r.status_code == 200
    data = r.json()
    assert data["reviewed"] is True
    assert data["review_note"] == "Checked - OK"
    # verify persistence
    r = requests.get(f"{API}/jobs/{demo_job['id']}/results", timeout=30)
    row = next(x for x in r.json() if x["id"] == rid)
    assert row["reviewed"] is True
    assert row["review_note"] == "Checked - OK"


# ---------- Reports ----------
@pytest.mark.parametrize("kind", ["exceptions", "full"])
def test_report_download(demo_job, kind):
    r = requests.get(f"{API}/jobs/{demo_job['id']}/report", params={"kind": kind}, timeout=60)
    assert r.status_code == 200
    assert "spreadsheet" in r.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(r.content))
    assert "Summary" in wb.sheetnames
    assert "Exceptions" in wb.sheetnames
    if kind == "full":
        assert "All Records" in wb.sheetnames


# ---------- Mapping validation ----------
def test_reconcile_missing_required_mapping(demo_job):
    sap_map = dict(demo_job["sap"]["suggested_mapping"])
    sap_map["gstin"] = None
    body = {
        "sap_mapping": sap_map,
        "cmp_mapping": demo_job["comparison"]["suggested_mapping"],
        "tolerance": 1.0, "sap_date_format": "AUTO", "cmp_date_format": "AUTO",
    }
    r = requests.post(f"{API}/jobs/{demo_job['id']}/reconcile", json=body, timeout=30)
    assert r.status_code == 400
    assert "GSTIN" in r.text or "Supplier GSTIN" in r.text


def test_reconcile_missing_tax_components(demo_job):
    cmp_map = dict(demo_job["comparison"]["suggested_mapping"])
    cmp_map["igst"] = None
    cmp_map["cgst"] = None
    cmp_map["sgst"] = None
    cmp_map["tax_amount"] = None
    body = {
        "sap_mapping": demo_job["sap"]["suggested_mapping"],
        "cmp_mapping": cmp_map,
        "tolerance": 1.0, "sap_date_format": "AUTO", "cmp_date_format": "AUTO",
    }
    r = requests.post(f"{API}/jobs/{demo_job['id']}/reconcile", json=body, timeout=30)
    assert r.status_code == 400


# ---------- Manual upload ----------
def test_manual_upload_csv_xlsx():
    sap = requests.get(f"{API}/demo/files/sap_csv", timeout=30).content
    gst = requests.get(f"{API}/demo/files/gstr2b_xlsx", timeout=30).content
    files = {
        "sap_file": ("sap.csv", sap, "text/csv"),
        "comparison_file": ("gstr.xlsx", gst,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }
    r = requests.post(f"{API}/jobs", files=files, data={"comparison_type": "gstr2b"}, timeout=60)
    assert r.status_code == 200, r.text
    job = r.json()
    created_jobs.append(job["id"])
    assert job["comparison"]["sheet"] == "B2B"
    assert job["comparison"]["header_row"] == 4


def test_manual_upload_purchase_register():
    sap = requests.get(f"{API}/demo/files/sap_csv", timeout=30).content
    gst = requests.get(f"{API}/demo/files/gstr2b_csv", timeout=30).content
    files = {
        "sap_file": ("sap.csv", sap, "text/csv"),
        "comparison_file": ("pr.csv", gst, "text/csv"),
    }
    r = requests.post(f"{API}/jobs", files=files, data={"comparison_type": "purchase_register"}, timeout=60)
    assert r.status_code == 200
    created_jobs.append(r.json()["id"])


# ---------- Error handling ----------
def test_upload_invalid_type():
    files = {
        "sap_file": ("x.pdf", b"%PDF-1.4 garbage", "application/pdf"),
        "comparison_file": ("y.csv", b"a,b\n1,2\n", "text/csv"),
    }
    r = requests.post(f"{API}/jobs", files=files, data={"comparison_type": "gstr2b"}, timeout=30)
    assert r.status_code == 400
    assert "Unsupported" in r.text or "CSV" in r.text


def test_upload_empty_file():
    files = {
        "sap_file": ("empty.csv", b"", "text/csv"),
        "comparison_file": ("y.csv", b"a,b\n1,2\n", "text/csv"),
    }
    r = requests.post(f"{API}/jobs", files=files, data={"comparison_type": "gstr2b"}, timeout=30)
    assert r.status_code == 400


def test_upload_bad_comparison_type():
    files = {
        "sap_file": ("x.csv", b"a,b\n1,2\n", "text/csv"),
        "comparison_file": ("y.csv", b"a,b\n1,2\n", "text/csv"),
    }
    r = requests.post(f"{API}/jobs", files=files, data={"comparison_type": "foo"}, timeout=30)
    assert r.status_code == 400


# ---------- Templates ----------
def test_templates_crud():
    body = {"name": "TEST_tpl_sap", "side": "sap", "mapping": {"gstin": "LicTradNum", "invoice_no": "NumAtCard"}}
    r = requests.post(f"{API}/templates", json=body, timeout=15)
    assert r.status_code == 200
    tpl = r.json()
    created_templates.append(tpl["id"])
    assert tpl["name"] == "TEST_tpl_sap"
    assert tpl["side"] == "sap"

    r = requests.get(f"{API}/templates", timeout=15)
    assert r.status_code == 200
    assert any(t["id"] == tpl["id"] for t in r.json())

    r = requests.get(f"{API}/templates", params={"side": "sap"}, timeout=15)
    assert all(t["side"] == "sap" for t in r.json())

    r = requests.delete(f"{API}/templates/{tpl['id']}", timeout=15)
    assert r.status_code == 200
    created_templates.remove(tpl["id"])


# ---------- Jobs CRUD ----------
def test_list_jobs(demo_job):
    r = requests.get(f"{API}/jobs", timeout=15)
    assert r.status_code == 200
    jobs = r.json()
    assert any(j["id"] == demo_job["id"] for j in jobs)


def test_get_nonexistent_job():
    r = requests.get(f"{API}/jobs/nonexistent-id-xxx", timeout=15)
    assert r.status_code == 404


# ---------- Pilot layouts ----------
PILOT_EXPECTED = {
    "s4_gst_register": {
        "ext": ".xlsx",
        "sap_date_format": "AUTO",
        "absolute": False,
        "matched": 27,
        "map": {"gstin": "Vendor GSTIN No.", "invoice_no": "Vendor Invoice No",
                "invoice_date": "Invoice Dt", "taxable_value": "Assessable Value",
                "igst": "IGST Amt", "cgst": "CGST Amt", "sgst": "SGST Amt"},
    },
    "b1_query_mdy": {
        "ext": ".csv",
        "sap_date_format": "MDY",
        "absolute": False,
        "matched": 27,
        "map": {"gstin": "GSTIN", "invoice_no": "Vendor Ref. No.",
                "invoice_date": "Document Date", "taxable_value": "Total Before Tax",
                "tax_amount": "Tax Amount (LC)"},
    },
    "ecc_technical": {
        "ext": ".csv",
        "sap_date_format": "AUTO",
        "absolute": True,
        "matched": 27,
        "map": {"gstin": "STCD3", "invoice_no": "XBLNR", "invoice_date": "BLDAT",
                "taxable_value": "HWBAS", "tax_amount": "WMWST"},
    },
    "s4_fiori": {
        "ext": ".xlsx",
        "sap_date_format": "AUTO",
        "absolute": False,
        "matched": 27,
        "map": {"gstin": "Tax Number 3", "invoice_no": "Reference",
                "invoice_date": "Invoice Date", "taxable_value": "Net Amount",
                "tax_amount": "Tax Amount"},
    },
}


def test_schema_pilot_layouts():
    r = requests.get(f"{API}/schema", timeout=15)
    data = r.json()
    assert "pilot_layouts" in data
    for k in PILOT_EXPECTED:
        assert k in data["pilot_layouts"]


@pytest.mark.parametrize("kind", list(PILOT_EXPECTED.keys()))
def test_pilot_layout_download(kind):
    r = requests.get(f"{API}/demo/files/{kind}", timeout=30)
    assert r.status_code == 200
    assert len(r.content) > 200


@pytest.mark.parametrize("kind", list(PILOT_EXPECTED.keys()))
def test_pilot_auto_mapping_and_reconcile(kind):
    exp = PILOT_EXPECTED[kind]
    sap = requests.get(f"{API}/demo/files/{kind}", timeout=30).content
    gst = requests.get(f"{API}/demo/files/gstr2b_xlsx", timeout=30).content
    files = {
        "sap_file": (f"pilot_{kind}{exp['ext']}", sap,
                     "text/csv" if exp["ext"] == ".csv" else
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "comparison_file": ("gstr.xlsx", gst,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }
    r = requests.post(f"{API}/jobs", files=files, data={"comparison_type": "gstr2b"}, timeout=60)
    assert r.status_code == 200, r.text
    job = r.json()
    created_jobs.append(job["id"])

    # header_row expectation for s4_gst_register = 4 (2 title rows + 1 blank + header)
    if kind == "s4_gst_register":
        assert job["sap"]["header_row"] == 4, f"header_row={job['sap']['header_row']}"

    sap_map = job["sap"]["suggested_mapping"]
    for k, expected_col in exp["map"].items():
        assert sap_map.get(k) == expected_col, \
            f"[{kind}] field {k}: expected {expected_col!r}, got {sap_map.get(k)!r}"

    # Reconcile with recommended settings
    body = {
        "sap_mapping": sap_map,
        "cmp_mapping": job["comparison"]["suggested_mapping"],
        "tolerance": 1.0,
        "sap_date_format": exp["sap_date_format"],
        "cmp_date_format": "AUTO",
        "absolute_amounts": exp["absolute"],
    }
    r = requests.post(f"{API}/jobs/{job['id']}/reconcile", json=body, timeout=60)
    assert r.status_code == 200, r.text
    s = r.json()["summary"]
    assert s["matched"] == exp["matched"], \
        f"[{kind}] expected matched={exp['matched']}, got {s['matched']}"


def test_pilot_ecc_without_absolute_flags_sign():
    """ECC trailing-minus with absolute_amounts OFF => 0 matched, reasons mention sign."""
    sap = requests.get(f"{API}/demo/files/ecc_technical", timeout=30).content
    gst = requests.get(f"{API}/demo/files/gstr2b_xlsx", timeout=30).content
    files = {
        "sap_file": ("ecc.csv", sap, "text/csv"),
        "comparison_file": ("gstr.xlsx", gst,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }
    r = requests.post(f"{API}/jobs", files=files, data={"comparison_type": "gstr2b"}, timeout=60)
    assert r.status_code == 200
    job = r.json()
    created_jobs.append(job["id"])
    body = {
        "sap_mapping": job["sap"]["suggested_mapping"],
        "cmp_mapping": job["comparison"]["suggested_mapping"],
        "tolerance": 1.0,
        "sap_date_format": "AUTO",
        "cmp_date_format": "AUTO",
        "absolute_amounts": False,
    }
    r = requests.post(f"{API}/jobs/{job['id']}/reconcile", json=body, timeout=60)
    assert r.status_code == 200
    s = r.json()["summary"]
    assert s["matched"] == 0, f"expected 0 matched (sign flip), got {s['matched']}"
    # Look at a results row for the sign hint
    r = requests.get(f"{API}/jobs/{job['id']}/results", timeout=30)
    rows = r.json()
    joined = " ".join(str(row.get("reason", "")) for row in rows).lower()
    assert "sign" in joined and "ignore amount sign" in joined, \
        f"amount-mismatch reason should mention sign / Ignore amount sign; got: {joined[:300]}"


def test_delete_job():
    r = requests.post(f"{API}/jobs/demo", timeout=60)
    jid = r.json()["id"]
    r = requests.delete(f"{API}/jobs/{jid}", timeout=15)
    assert r.status_code == 200
    r = requests.get(f"{API}/jobs/{jid}", timeout=15)
    assert r.status_code == 404
